"""
文档引擎单元测试
"""

import os
import sys
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import tests._base  # noqa: F401

import asyncio
import io
from docx import Document
from backend.app.services.document_engine import DocumentEngine
from backend.app.services.template_engine import TemplateEngine
from backend.app.database import db


class DocumentEngineTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        async def _init():
            await db.init()
            await TemplateEngine.initialize()

        asyncio.run(_init())

    def _sample_sections(self):
        return [
            {"type": "body", "title": "第一章 绪论", "content": "这是绪论内容。\n## 1.1 研究背景\n背景描述文本。", "order": 0},
            {"type": "body", "title": "第二章 方法", "content": "方法章节正文内容。", "order": 1},
        ]

    def test_generate_docx_bytes(self):
        """生成 docx 应返回有效字节，且可被重新打开"""
        doc = DocumentEngine.create_document(
            sections=self._sample_sections(),
            template_id="default",
            title="测试论文",
            abstract={"chinese": "这是摘要内容。", "keywords_zh": ["测试", "论文"]},
        )
        data = DocumentEngine.to_bytes(doc)
        self.assertGreater(len(data), 1000, "docx 字节数应足够大")

        # 重新打开验证结构
        reopened = Document(io.BytesIO(data))
        texts = [p.text for p in reopened.paragraphs]
        joined = "\n".join(texts)
        self.assertIn("测试论文", joined, "标题应存在")
        self.assertIn("这是摘要内容", joined, "摘要应存在")
        self.assertIn("这是绪论内容", joined, "正文应存在")
        self.assertIn("研究背景", joined, "小节标题应存在")

    def test_template_applied_to_document(self):
        """模板样式应应用到生成的文档"""
        doc = DocumentEngine.create_document(
            sections=self._sample_sections(), template_id="default", title="样式测试"
        )
        normal = doc.styles["Normal"]
        self.assertEqual(normal.font.size.pt, 12, "正文字号应为12pt")
        self.assertEqual(normal.paragraph_format.line_spacing, 1.5, "行距应为1.5")

    def test_markdown_heading_parsed(self):
        """Markdown 标题应被解析为文档标题"""
        sections = [
            {"type": "body", "title": "", "content": "## 二级标题\n正文内容。", "order": 0}
        ]
        doc = DocumentEngine.create_document(sections=sections, template_id="default")
        joined = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("二级标题", joined)
        self.assertIn("正文内容", joined)

    def test_deep_markdown_heading_no_hash(self):
        """4 级及以下标题（####）不应残留 # 符号"""
        sections = [
            {"type": "body", "title": "", "content": "#### 5.2.1 边界问题\n这不是标题的正文。", "order": 0}
        ]
        doc = DocumentEngine.create_document(sections=sections, template_id="default")
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("####", text, "不应残留 Markdown 井号")
        self.assertIn("5.2.1 边界问题", text, "标题文字应被保留")
        self.assertIn("这不是标题的正文", text)

    def test_toc_prefilled(self):
        """目录应预填章节标题，而非只有占位文本"""
        sections = [
            {"type": "body", "title": "第一章 绪论", "content": "## 1.1 研究背景\n内容。", "order": 0},
            {"type": "body", "title": "第二章 方法", "content": "正文。", "order": 1},
        ]
        doc = DocumentEngine.create_document(
            sections=sections, template_id="default", title="目录测试"
        )
        joined = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("第一章 绪论", joined, "目录应包含一级标题")
        self.assertIn("1.1 研究背景", joined, "目录应包含二级标题")
        self.assertIn("第二章 方法", joined)

    def test_normalize_reference(self):
        """参考文献条目应被规范化清洗"""
        raw = "张三, 李四. 深度学习综述   [J]. 计算机学报,2020,43(1):1-20"
        out = DocumentEngine.normalize_reference(raw)
        self.assertTrue(out.endswith("."), "应补结尾句点")
        self.assertNotIn("  ", out, "不应有多余连续空格")
        self.assertIn("[J]", out, "应保留文献类型")
        self.assertIn("计算机学报, 2020, 43(1): 1-20.", out, "年份/页码空格应规范")

        # 无文献类型时应补 [J]
        raw2 = "王五. 无类型标注的条目"
        out2 = DocumentEngine.normalize_reference(raw2)
        self.assertIn("[J]", out2, "应补文献类型")
        self.assertTrue(out2.endswith("."))

        # 行首编号应被剥离（导出时统一重新编号）
        raw3 = "[3] 赵六. 另一篇文献[M]. 北京: 某出版社, 2021"
        out3 = DocumentEngine.normalize_reference(raw3)
        self.assertFalse(out3.startswith("[3]"), "行首编号应剥离")
        self.assertIn("[M]", out3)


if __name__ == "__main__":
    unittest.main(verbosity=2)